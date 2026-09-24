from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"C:\Users\chuak\Downloads\Test-Case-Template.docx")
OUTPUT = Path(r"C:\Users\chuak\Desktop\menrosystem\MENRO-Test-Cases-Actual-Results.docx")


cases = [
    ["TC-SYS-001", "Frontend production build completes successfully", "REQ-SYS-01", "Frontend Build", "System", "H", "Major", "Frontend dependencies installed", "Current menrosystem source", "1) Run the production build.\n2) Wait for bundling.\n3) Review the exit result.", "All modules compile and production assets are generated without build errors.", "PASSED - 2,497 modules transformed; build exited successfully.", ""],
    ["TC-SYS-002", "Frontend source passes code-quality validation", "REQ-SYS-02", "Frontend Lint", "Unit", "M", "Minor", "ESLint configuration and dependencies available", "Current src files", "1) Run the lint command.\n2) Review reported errors and exit code.", "Lint completes with no errors.", "PASSED - ESLint completed with exit code 0.", ""],
    ["TC-AUTH-003", "Profile uses the authenticated user's UID", "REQ-AUTH-01", "Authentication Profile", "Integration", "H", "Critical", "Authenticated user context is provided", "Authenticated UID and Firestore profile", "1) Request the user profile.\n2) Compare the queried UID with the authenticated UID.\n3) verify returned name.", "The system reads the profile belonging to the authenticated UID only.", "PASSED - automated test confirmed authenticated UID and stored name are used.", ""],
    ["TC-AUTH-004", "Participant can open only their own planting report", "REQ-AUTH-02", "Report Authorization", "Integration", "H", "Critical", "Two participant records and reports exist", "Participant A token; Participant B report ID", "1) request Participant A's report.\n2) request Participant B's report using Participant A identity.", "Own report is returned; another participant's report is denied.", "PASSED - ownership authorization test succeeded.", ""],
    ["TC-AUTH-005", "Only Staff can perform final planting-report decisions", "REQ-AUTH-03", "Role Authorization", "Integration", "H", "Critical", "Pending Review report exists", "Staff, Admin, and Participant roles", "1) call decision routes using each role.\n2) verify allowed and denied results.", "Staff is allowed; Admin and Participant cannot use Staff-only decision routes.", "PASSED - automated route-role test confirmed Staff-only access.", ""],
    ["TC-AUTH-006", "Staff and Admin can view event participants", "REQ-AUTH-04", "Event Authorization", "Integration", "H", "Major", "Event with participants exists", "Staff, Admin, and guest identities", "1) request event participants as Staff.\n2) repeat as Admin.\n3) attempt guest decision access.", "Staff and Admin can read participants; guests have no decision route.", "PASSED - role-access assertions succeeded.", ""],
    ["TC-AUTH-007", "Rejected report records reason and authenticated reviewer", "REQ-AUTH-05", "Audit Trail", "Integration", "H", "Major", "Report is ready for Staff review", "Blank reason; valid reason; Staff UID", "1) reject without a reason.\n2) reject with a reason.\n3) inspect reviewer fields.", "Blank reason is rejected; valid rejection stores the server-authenticated reviewer.", "PASSED - reason validation and reviewer persistence succeeded.", ""],
    ["TC-SEC-008", "Guest invitation is limited to its owner and event", "REQ-SEC-01", "Guest Invitation", "Integration", "H", "Critical", "Approved request and scheduled event exist", "Requester UID, invitation token, event ID", "1) retrieve the invitation as its requester.\n2) use it for the correct event.\n3) attempt mismatched access.", "Invitation access is owner-scoped and guest activity remains event-scoped.", "PASSED - connected-workflow security assertions succeeded.", ""],
    ["TC-SEC-009", "Missing photo GPS metadata is rejected", "REQ-SEC-02", "Planting Evidence", "Integration", "H", "Major", "Participant has a released planting allocation", "Valid image without EXIF GPS", "1) submit the image as planting evidence.\n2) inspect the response and saved records.", "The submission is rejected and no client coordinate can replace missing EXIF GPS.", "PASSED - missing-GPS evidence test succeeded.", ""],
    ["TC-SEC-010", "Client coordinates cannot replace missing EXIF GPS", "REQ-SEC-03", "Evidence Integrity", "Integration", "H", "Major", "Evidence image has no EXIF GPS", "Image plus client-supplied latitude and longitude", "1) attach client coordinates.\n2) submit the non-geotagged image.\n3) inspect validation result.", "Server still rejects the image because trusted photo metadata is absent.", "PASSED - client-coordinate substitution was rejected.", ""],
    ["TC-SEC-011", "Invalid image bytes are rejected", "REQ-SEC-04", "Upload Validation", "Unit", "H", "Major", "Evidence validation service is available", "File named as an image with invalid bytes", "1) pass invalid bytes to image verification.\n2) inspect validation result.", "The file is rejected as invalid image content.", "PASSED - invalid image-byte test succeeded.", ""],
    ["TC-SEC-012", "Screenshot detection requires reliable evidence", "REQ-SEC-05", "Image Verification", "Unit", "M", "Major", "Image metadata and detection signals are available", "Single weak signal and combined strong signals", "1) test one weak signal.\n2) test combined signals.\n3) compare classifications.", "A weak signal alone does not cause a false positive; strong combined evidence is flagged.", "PASSED - screenshot classification test succeeded.", ""],
    ["TC-SEC-013", "Photo location is checked against the registered site", "REQ-SEC-06", "Geofence Validation", "Unit", "H", "Major", "Registered site polygon exists", "Inside and outside GPS coordinates", "1) verify coordinates inside the polygon.\n2) verify coordinates outside it.", "Inside coordinates pass; outside coordinates are reported as a mismatch.", "PASSED - site-polygon verification test succeeded.", ""],
    ["TC-DATA-014", "Invalid seedling release leaves data unchanged", "REQ-DATA-01", "Inventory Transaction", "Integration", "H", "Critical", "Inventory, request, and event records exist", "Release quantity exceeding the allowed stock", "1) attempt the invalid release.\n2) inspect inventory, event allocation, and distribution records.", "The operation fails atomically and all related records remain unchanged.", "PASSED - rollback and unchanged-state assertions succeeded.", ""],
    ["TC-DATA-015", "Seedling request rejects a site in another barangay", "REQ-DATA-02", "Request Validation", "Integration", "H", "Major", "Participant barangay and real site record exist", "Site located in a different barangay", "1) create a request using the mismatched site.\n2) inspect response and stored requests.", "The request is rejected and no invalid request is stored.", "PASSED - barangay/site consistency test succeeded.", ""],
]


def set_cell_text_preserving_format(cell, text):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0

    template_rpr = None
    for run in paragraph.runs:
        if run._element.rPr is not None:
            template_rpr = deepcopy(run._element.rPr)
            break
    for run in list(paragraph.runs):
        paragraph._p.remove(run._element)
    run = paragraph.add_run(text)
    if template_rpr is not None:
        if run._element.rPr is not None:
            run._element.remove(run._element.rPr)
        run._element.insert(0, template_rpr)
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(7)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


doc = Document(SOURCE)

# Fill only the editable metadata slots while preserving the template.
set_cell_text_preserving_format(doc.tables[0].cell(0, 1), "MENRO Reforestation Management System")
set_cell_text_preserving_format(doc.tables[0].cell(2, 1), date.today().strftime("%B %d, %Y"))

table = doc.tables[1]
assert len(table.rows) == 17 and len(table.columns) == 13

# Row 0 is the header, row 1 is the supplied sample, rows 2-16 are numbered 1-15.
for row_index, values in enumerate(cases, start=2):
    row = table.rows[row_index]
    for col_index, value in enumerate(values):
        set_cell_text_preserving_format(row.cells[col_index], value)
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)

# Repeat the original header row if the table continues to later pages.
header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
if header_tr_pr.find(qn("w:tblHeader")) is None:
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_tr_pr.append(repeat)

doc.core_properties.title = "MENRO System Test Cases and Actual Results"
doc.save(OUTPUT)
print(OUTPUT)

